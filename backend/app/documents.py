import io
import os
import zipfile

import fitz
from docx import Document
from PIL import Image

MAX_BYTES = 10 * 1024 * 1024
MAX_PAGES = 20
Image.MAX_IMAGE_PIXELS = 20_000_000
TYPES = {".pdf": "application/pdf", ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document", ".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg"}


class ExtractionError(ValueError):
    pass


def ocr(image):
    import pytesseract
    if os.getenv("TESSERACT_CMD"):
        pytesseract.pytesseract.tesseract_cmd = os.environ["TESSERACT_CMD"]
    try:
        return pytesseract.image_to_string(image, timeout=20)
    except pytesseract.TesseractNotFoundError as exc:
        raise ExtractionError("OCR unavailable. Install Tesseract and set TESSERACT_CMD in backend/.env. Text-based PDFs still work.") from exc
    except RuntimeError as exc:
        raise ExtractionError("OCR timed out. Supply a smaller, clearer document.") from exc


def extract(data, extension):
    if extension not in TYPES:
        raise ExtractionError("Supported file types: PDF, DOCX, PNG, JPG.")
    if not data or len(data) > MAX_BYTES:
        raise ExtractionError("File must be nonempty and at most 10 MB.")
    sections, warnings = [], []
    try:
        if extension == ".pdf":
            if not data.startswith(b"%PDF-"):
                raise ExtractionError("File signature does not match PDF.")
            with fitz.open(stream=data, filetype="pdf") as pdf:
                if pdf.needs_pass or len(pdf) > MAX_PAGES:
                    raise ExtractionError("Use an unencrypted PDF with at most 20 pages.")
                for i, page in enumerate(pdf):
                    text = page.get_text()
                    if len(text.strip()) < 15:
                        if page.rect.width * page.rect.height * 2.25 > 20_000_000:
                            raise ExtractionError("Scanned page dimensions exceed processing limit.")
                        pixmap = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5))
                        if pixmap.width * pixmap.height > 20_000_000:
                            raise ExtractionError("Scanned page dimensions exceed processing limit.")
                        text = ocr(Image.open(io.BytesIO(pixmap.tobytes("png"))))
                        warnings.append(f"Page {i + 1}: OCR text requires manual review.")
                    sections.append({"location": f"page {i + 1}", "text": text[:30000]})
        elif extension == ".docx":
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if sum(entry.file_size for entry in archive.infolist()) > 30 * 1024 * 1024:
                    raise ExtractionError("DOCX expanded size exceeds 30 MB.")
                if "word/document.xml" not in archive.namelist():
                    raise ExtractionError("File is not a Word document.")
            doc = Document(io.BytesIO(data))
            sections = [{"location": f"paragraph {i + 1}", "text": p.text} for i, p in enumerate(doc.paragraphs) if p.text.strip()]
            for i, table in enumerate(doc.tables):
                sections.append({"location": f"table {i + 1}", "text": "\n".join(" | ".join(c.text for c in row.cells) for row in table.rows)})
        else:
            with Image.open(io.BytesIO(data)) as image:
                expected = "PNG" if extension == ".png" else "JPEG"
                if image.format != expected or image.width * image.height > 20_000_000:
                    raise ExtractionError("Image format or dimensions are invalid.")
                sections = [{"location": "image 1", "text": ocr(image)}]
                warnings.append("OCR text requires manual review; confidence is not calibrated.")
    except ExtractionError:
        raise
    except Exception as exc:
        raise ExtractionError("Document could not be read. Check its format, encryption and image quality.") from exc
    text = "\n".join(s["text"] for s in sections)
    if not text.strip():
        raise ExtractionError("No readable text was extracted. Supply a readable document.")
    if len(text) > 250000:
        raise ExtractionError("Extracted text exceeds 250,000 characters.")
    return {"extracted_text": text, "sections": sections, "warnings": warnings, "processing_status": "ready"}